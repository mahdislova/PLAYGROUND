"""Editable 6x9 book-layout DOCX for THE PORT RATS, built from book.json."""
import json, os, re, html, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

H = os.path.dirname(os.path.abspath(__file__))
book = json.load(open(os.path.join(H, "book.json"), encoding="utf-8"))
FONT = "Garamond"
BROWN = RGBColor(0x6B, 0x56, 0x41)
LINE = Pt(14.6)

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(6), Inches(9)
sec.top_margin, sec.bottom_margin = Inches(0.72), Inches(0.8)
sec.left_margin, sec.right_margin = Inches(0.86), Inches(0.6)   # inside / outside with mirror margins
sec.footer_distance = Inches(0.4)

st = d.settings.element
for tag, val in (("w:mirrorMargins", None), ("w:autoHyphenation", None),
                 ("w:consecutiveHyphenLimit", "2"), ("w:hyphenationZone", "360")):
    el = OxmlElement(tag)
    if val:
        el.set(qn("w:val"), val)
    st.append(el)


def set_font(style_or_run, name=FONT):
    style_or_run.font.name = name
    rpr = style_or_run.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


S = {s.name: s for s in d.styles}
normal = S["Normal"]
set_font(normal)
normal.font.size = Pt(11)
pf = normal.paragraph_format
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.first_line_indent = Inches(0.2)
pf.space_before = pf.space_after = Pt(0)
pf.line_spacing = LINE
pf.widow_control = True


def style(name, base=normal, **kw):
    s = d.styles.add_style(name, 1)
    s.base_style = base
    set_font(s)
    f = s.paragraph_format
    for k, v in kw.items():
        setattr(f, k, v)
    return s


def ctr(name, **kw):
    return style(name, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=Inches(0), **kw)


s_first = style("First Paragraph", first_line_indent=Inches(0))
s_cnum = ctr("Chapter Number", space_before=Inches(1.1), space_after=Pt(4), keep_with_next=True, page_break_before=True)
s_cnum.font.italic = True; s_cnum.font.size = Pt(12); s_cnum.font.color.rgb = BROWN
s_ctitle = ctr("Chapter Title", space_after=Pt(8), keep_with_next=True)
s_ctitle.font.size = Pt(22); s_ctitle.font.small_caps = True
s_orn = ctr("Ornament", space_before=Pt(9), space_after=Pt(9), keep_with_next=True)
s_head_orn = ctr("Chapter Ornament", space_after=Pt(22), keep_with_next=True)
s_center = ctr("Centered")
s_copy = style("Copyright Text", first_line_indent=Inches(0), alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=Pt(8))
s_copy.font.size = Pt(9)


def add_runs(p, frag, small_caps=False):
    """Minimal HTML -> runs: <em>, <strong>, entities."""
    it, bd = False, False
    for tok in re.split(r"(<[^>]+>)", frag):
        if not tok:
            continue
        if tok.startswith("<"):
            t = tok.lower()
            if t.startswith("<em"): it = True
            elif t.startswith("</em"): it = False
            elif t.startswith("<strong"): bd = True
            elif t.startswith("</strong"): bd = False
            continue
        r = p.add_run(html.unescape(tok))
        r.italic = it or None
        r.bold = bd or None
        if small_caps:
            r.font.small_caps = True
            r.font.spacing = Pt(0.6)


def image_para(path, width, sty=s_orn):
    p = d.add_paragraph(style=sty)
    p.add_run().add_picture(path, width=width)
    return p


def drop_cap(letter, quote=""):
    """Word-native 3-line dropped initial (framePr)."""
    p = d.add_paragraph(style=s_first)
    ppr = p._p.get_or_add_pPr()
    fr = OxmlElement("w:framePr")
    for k, v in (("w:dropCap", "drop"), ("w:lines", "3"), ("w:wrap", "around"),
                 ("w:vAnchor", "text"), ("w:hAnchor", "text")):
        fr.set(qn(k), v)
    ppr.insert(0, fr)
    p.paragraph_format.line_spacing = Pt(44)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(quote + letter)
    r.font.size = Pt(52)
    r.font.color.rgb = RGBColor(0x3A, 0x2A, 0x1C)
    rpr = r._r.get_or_add_rPr()
    pos = OxmlElement("w:position"); pos.set(qn("w:val"), "-6"); rpr.append(pos)
    return p


def page_break():
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


IMG = lambda n: os.path.join(H, "img", n + ".png")

# ---------------- title page ----------------
p = d.add_paragraph(style=s_center); p.paragraph_format.space_before = Inches(0.9)
p.add_run().add_picture(IMG("emblem"), width=Inches(3.5))
p = d.add_paragraph(style=s_center); p.paragraph_format.space_before = Pt(26)
r = p.add_run("THE"); r.font.size = Pt(13); r.font.spacing = Pt(9); r.font.color.rgb = BROWN
p = d.add_paragraph(style=s_center); p.paragraph_format.line_spacing = Pt(46)
r = p.add_run("Port Rats"); r.font.size = Pt(40); r.font.small_caps = True; r.font.spacing = Pt(3)
p = d.add_paragraph(style=s_center); p.paragraph_format.space_before = Pt(6)
r = p.add_run("a novel of the Lowers"); r.italic = True; r.font.size = Pt(13); r.font.color.rgb = BROWN
p = d.add_paragraph(style=s_center); p.paragraph_format.space_before = Pt(18)
p.add_run().add_picture(IMG("ashleaf"), width=Inches(0.75))
p = d.add_paragraph(style=s_center); p.paragraph_format.space_before = Inches(1.2)
r = p.add_run("PATRICK"); r.font.size = Pt(15); r.font.spacing = Pt(6)
page_break()

# ---------------- copyright ----------------
first = True
for line in ["This is a work of fiction. Names, characters, places, and incidents are the product of the author’s imagination. Any resemblance to actual persons, living or dead, or to actual events, is coincidental.",
             "All characters depicted are adults.",
             "For mature readers. Contains explicit sexual content, violence, and strong language.",
             "Copyright © 2026 by Patrick. All rights reserved.",
             "Draft edition · October 2026",
             "Set in Garamond."]:
    p = d.add_paragraph(line, style=s_copy)
    if first:
        p.paragraph_format.space_before = Inches(3.6); first = False
p = d.add_paragraph(style=s_copy); p.add_run().add_picture(IMG("mooring"), width=Inches(1.0))
page_break()

# ---------------- contents ----------------
p = d.add_paragraph(style=s_center); p.paragraph_format.space_before = Inches(0.6)
r = p.add_run("Contents"); r.font.size = Pt(20); r.font.small_caps = True; r.font.spacing = Pt(2)
image_para(IMG("ashleaf"), Inches(0.6), s_center).paragraph_format.space_after = Pt(18)
for c in book:
    p = d.add_paragraph(style=s_first)
    p.paragraph_format.left_indent = Inches(0.9)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.85))
    r = p.add_run(c["num"]); r.italic = True; r.font.color.rgb = BROWN
    p.add_run("\t")
    r = p.add_run(c["title"]); r.font.small_caps = True; r.font.spacing = Pt(0.5)

# ---------------- body: new section, folios start at 1 ----------------
body = d.add_section(WD_SECTION.ODD_PAGE)
body.footer.is_linked_to_previous = False
pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); body._sectPr.append(pg)
fp = body.footer.paragraphs[0]; fp.style = s_center
fr = fp.add_run()
for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if kind:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
    fr._r.append(e)
fr.font.size = Pt(9.5)

for ci, c in enumerate(book):
    p = d.add_paragraph(f"Chapter {c['num']}", style=s_cnum)
    if ci == 0:
        p.paragraph_format.page_break_before = False   # section break already starts the page
    d.add_paragraph(c["title"], style=s_ctitle)
    image_para(IMG("ratline"), Inches(2.3), s_head_orn)
    for b in c["blocks"]:
        if b["t"] == "scene":
            image_para(IMG("mooring"), Inches(1.05))
        elif b["t"] == "opener":
            ini = b["initial"]
            m = re.match(r"^([“‘\"']+)(.*)$", ini)
            q, letter = (m.group(1), m.group(2)) if m else ("", ini)
            drop_cap(letter, q)
            p = d.add_paragraph(style=s_first)
            add_runs(p, b["lead"], small_caps=True)
            add_runs(p, b["tail"])
        else:
            p = d.add_paragraph(style=s_first if b["t"] == "flush" else normal)
            add_runs(p, b["html"])

image_para(IMG("ashleaf"), Inches(0.8), s_center).paragraph_format.space_before = Pt(28)
p = d.add_paragraph("End of the draft · October 2026", style=s_center)
p.runs[0].italic = True; p.runs[0].font.color.rgb = BROWN; p.runs[0].font.size = Pt(10)

d.core_properties.title = "The Port Rats"
d.core_properties.author = "Patrick"
out = os.path.join(H, "The Port Rats - Editable 6x9.docx")
d.save(out)
print("saved", out)
